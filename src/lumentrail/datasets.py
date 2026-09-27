"""Explicit dataset downloads and small, original offline fixtures."""

from __future__ import annotations

import json
import csv
import hashlib
import urllib.request
import xml.etree.ElementTree as ElementTree
import zipfile
from pathlib import Path

MEDQUAD_ARCHIVE = "https://github.com/abachaa/MedQuAD/archive/refs/heads/master.zip"
HEALTHSEARCHQA_WORKBOOK = (
    "https://static-content.springer.com/esm/art%3A10.1038%2Fs41586-023-06291-2/"
    "MediaObjects/41586_2023_6291_MOESM6_ESM.xlsx"
)

DEMO_PASSAGES = [
    {"id": "demo-asthma", "question": "What is asthma?", "answer": "Asthma is a long-term condition that can make breathing difficult.", "source": "synthetic-demo", "entities": ["asthma"]},
    {"id": "demo-allergy", "question": "What is an allergy?", "answer": "An allergy is an immune reaction to a substance that is usually harmless to others.", "source": "synthetic-demo", "entities": ["allergy"]},
    {"id": "demo-af", "question": "What is atrial fibrillation?", "answer": "Atrial fibrillation, often shortened to AF, is an irregular heart rhythm.", "source": "synthetic-demo", "entities": ["atrial fibrillation", "AF"]},
    {"id": "demo-eczema", "question": "What is eczema?", "answer": "Eczema is a group of conditions that can cause itchy, inflamed skin.", "source": "synthetic-demo", "entities": ["eczema"]},
    {"id": "demo-citation", "question": "How should a health answer use sources?", "answer": "Check that each claim is supported by the cited passage and show when the source was updated.", "source": "synthetic-demo", "entities": ["citation"]},
]

DEMO_JUDGEMENTS = [
    {"question": "What is asthma?", "relevant_ids": ["demo-asthma"]},
    {"question": "What is AF?", "relevant_ids": ["demo-af"]},
    {"question": "How do I cite health evidence?", "relevant_ids": ["demo-citation"]},
]


def write_jsonl(path: Path, records: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as output:
        for record in records:
            output.write(json.dumps(record, ensure_ascii=False) + "\n")


def read_jsonl(path: Path) -> list[dict]:
    with path.open(encoding="utf-8") as source:
        return [json.loads(line) for line in source if line.strip()]


def csv_passages(path: Path, question_column: str, answer_column: str, source_url: str) -> list[dict]:
    """Adapt a downloaded Kaggle or Hugging Face CSV without assuming its schema."""
    records = []
    seen: set[str] = set()
    with path.open(encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        if not reader.fieldnames or not {question_column, answer_column}.issubset(reader.fieldnames):
            raise ValueError(f"CSV must contain {question_column!r} and {answer_column!r} columns")
        for row in reader:
            question = (row.get(question_column) or "").strip()
            answer = (row.get(answer_column) or "").strip()
            if not question or not answer:
                continue
            passage_id = hashlib.sha256(f"{question}\n{answer}".encode()).hexdigest()[:20]
            if passage_id in seen:
                continue
            seen.add(passage_id)
            records.append({"id": f"csv:{passage_id}", "question": question,
                            "answer": answer, "source": source_url, "entities": []})
    return records


def download(url: str, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    request = urllib.request.Request(url, headers={"User-Agent": "LumenTrail/0.1"})
    with urllib.request.urlopen(request, timeout=60) as response, destination.open("wb") as output:
        while block := response.read(1024 * 1024):
            output.write(block)


def medquad_records(archive: Path, limit: int | None = None) -> list[dict]:
    records: list[dict] = []
    with zipfile.ZipFile(archive) as source:
        for name in sorted(source.namelist()):
            if not name.lower().endswith(".xml") or "/__macosx/" in name.lower():
                continue
            try:
                root = ElementTree.fromstring(source.read(name))
            except ElementTree.ParseError:
                continue
            document = root if root.tag == "Document" else root.find(".//Document")
            source_url = document.get("url", "") if document is not None else ""
            for position, pair in enumerate(root.findall(".//QAPair")):
                question = " ".join((pair.findtext("Question") or "").split())
                answer = " ".join((pair.findtext("Answer") or "").split())
                # The source intentionally omits answers from three copyrighted subsets.
                if not question or not answer:
                    continue
                source_file = "/".join(Path(name).parts[1:])
                records.append({
                    "id": f"medquad:{source_file}:{position}",
                    "question": question,
                    "answer": answer,
                    "source": source_url or f"https://github.com/abachaa/MedQuAD/blob/master/{source_file}",
                    "entities": [],
                })
                if limit and len(records) >= limit:
                    return records
    return records


def healthsearchqa_records(workbook: Path, limit: int | None = None) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("Install workbook support: pip install -e '.[healthsearchqa]'") from error
    records: list[dict] = []
    seen: set[str] = set()
    source = load_workbook(workbook, read_only=True, data_only=True)
    try:
        for row in source.active.iter_rows(values_only=True):
            for cell in row:
                if not isinstance(cell, str):
                    continue
                question = " ".join(cell.split())
                if not question.endswith("?") or question.lower() in seen:
                    continue
                seen.add(question.lower())
                records.append({"question": question, "relevant_ids": [], "needs_review": True})
                if limit and len(records) >= limit:
                    return records
    finally:
        source.close()
    return records
